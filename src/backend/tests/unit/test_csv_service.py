import io
import json
from datetime import datetime
from types import SimpleNamespace

import openpyxl
import pytest

from app.copilot.config import CopilotConfig
from app.models import CsvColumnMapping
from app.services.csv_service import (
    BANK_EXPORT_COLUMNS,
    ColumnMappingError,
    CsvImportService,
    _parse_italian_number,
    _parse_mapped_datetime,
    _prepare_fineco_csv_content,
    _read_mapping_table,
    _read_xlsx_rows,
    _side_from_value,
)


def test_prepare_fineco_csv_content_keeps_embedded_header_row():
    header = ";".join(
        [
            "Operazione",
            "Data valuta",
            "Descrizione",
            "Titolo",
            "Isin",
            "Segno",
            "Quantita",
            "Divisa",
            "Prezzo",
            "Cambio",
            "Controvalore",
            "Commissioni Fondi Sw/Ingr/Uscita",
            "Commissioni Fondi Banca Corrispondente",
            "Spese Fondi Sgr",
            "Commissioni amministrato",
        ]
    )
    content = "\n".join(
        [
            "estratto conto fineco",
            "cliente abc",
            header,
            "01/03/2026;04/03/2026;ACQUISTO;ETF TEST;IE00TEST0001;A;10;EUR;100,00;1;1000,00;0;0;0;2,95",
        ]
    )

    normalized = _prepare_fineco_csv_content(content, skip_rows=7)

    assert normalized.splitlines()[0].lower().startswith("operazione;data valuta;descrizione")


def test_prepare_fineco_csv_content_adds_header_for_headerless_export():
    content = "\n".join(
        [
            "metadata 1",
            "metadata 2",
            "metadata 3",
            "metadata 4",
            "metadata 5",
            "metadata 6",
            "metadata 7",
            "01/03/2026;04/03/2026;ACQUISTO;ETF TEST;IE00TEST0001;A;10;EUR;100,00;1;1000,00;0;0;0;2,95",
        ]
    )

    normalized = _prepare_fineco_csv_content(content, skip_rows=7)

    assert normalized.splitlines()[0] == ";".join(BANK_EXPORT_COLUMNS)


def test_read_xlsx_rows_detects_fineco_header_without_fixed_skip():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["report fineco"])
    worksheet.append([""])
    worksheet.append(
        [
            "Operazione",
            "Data valuta",
            "Descrizione",
            "Titolo",
            "Isin",
            "Segno",
            "Quantita",
            "Divisa",
            "Prezzo",
            "Cambio",
            "Controvalore",
            "Commissioni Fondi Sw/Ingr/Uscita",
            "Commissioni Fondi Banca Corrispondente",
            "Spese Fondi Sgr",
            "Commissioni amministrato",
        ]
    )
    worksheet.append(
        [
            "01/03/2026",
            "04/03/2026",
            "ACQUISTO",
            "ETF TEST",
            "IE00TEST0001",
            "A",
            "10",
            "EUR",
            "100,00",
            "1",
            "1000,00",
            "0",
            "0",
            "0",
            "2,95",
        ]
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    rows = _read_xlsx_rows(buffer.getvalue(), skip_rows=7, broker="fineco")

    assert len(rows) == 1
    assert rows[0]["Operazione"] == "01/03/2026"


REDUCED_FINECO_HEADER = [
    "Operazione",
    "Data valuta",
    "Descrizione",
    "Titolo",
    "Isin",
    "Segno",
    "Quantita",
    "Divisa",
    "Prezzo",
    "Cambio",
    "Controvalore",
    "Commissioni amministrato",
]


def test_read_xlsx_rows_detects_fineco_header_without_fund_fee_columns():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Dossier n.: 00000000"])
    worksheet.append(["Intestazione Dossier: MARIO ROSSI"])
    worksheet.append([""])
    worksheet.append(["RISULTATO RICERCA MOVIMENTI TITOLI"])
    worksheet.append([""])
    worksheet.append(REDUCED_FINECO_HEADER)
    worksheet.append([])
    worksheet.append(
        ["13/10/2026", "15/10/2026", "Compravendita titoli", "ISHS EUR CO 1-5 EUR", "IE00B4L60045", "A", 375, "EUR", 105.23067, 1, 39461.5, 19]
    )
    worksheet.append(
        ["13/10/2026", "15/10/2026", "Compravendita titoli", "VAN FTSE ALL USD-AC", "IE00BK5BQT80", "V", 1, "EUR", 153.89, 1, 153.89, 19]
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    rows = _read_xlsx_rows(buffer.getvalue(), skip_rows=7, broker="fineco")

    assert len(rows) == 2
    assert rows[0]["Isin"] == "IE00B4L60045"
    assert rows[1]["Isin"] == "IE00BK5BQT80"


def test_read_xlsx_rows_keeps_three_decimal_numeric_cells_as_decimals():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(REDUCED_FINECO_HEADER)
    worksheet.append(
        ["01/03/2026", "04/03/2026", "Compravendita titoli", "ETF TEST", "IE00TEST0001", "A", 10, "EUR", 105.123, 1, 1051.23, 2.5]
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    rows = _read_xlsx_rows(buffer.getvalue(), broker="fineco")

    assert _parse_italian_number(rows[0]["Prezzo"]) == 105.123
    assert _parse_italian_number(rows[0]["Quantita"]) == 10
    assert _parse_italian_number(rows[0]["Controvalore"]) == 1051.23


def test_prepare_fineco_csv_content_detects_header_without_fund_fee_columns():
    content = "\n".join(
        [
            "Dossier n.: 00000000",
            "RISULTATO RICERCA MOVIMENTI TITOLI",
            ";".join(REDUCED_FINECO_HEADER),
            "13/10/2026;15/10/2026;Compravendita titoli;ETF;IE00B4L60045;A;375;EUR;105,23;1;39461,5;19",
        ]
    )

    normalized = _prepare_fineco_csv_content(content, skip_rows=7)

    assert normalized.splitlines()[0] == ";".join(REDUCED_FINECO_HEADER)
    assert len(normalized.splitlines()) == 2


class _FakeRepo:
    pass


def test_parse_public_portfolio_file_aggregates_fineco_rows_into_positions():
    service = CsvImportService(_FakeRepo())
    content = "\n".join(
        [
            "metadata 1",
            "metadata 2",
            "metadata 3",
            "metadata 4",
            "metadata 5",
            "metadata 6",
            "metadata 7",
            "01/03/2026;04/03/2026;ACQUISTO;ETF TEST;IE00TEST0001;A;10;EUR;100,00;1;1000,00;0;0;0;2,95",
            "05/03/2026;06/03/2026;ACQUISTO;ETF TEST;IE00TEST0001;A;5;EUR;100,00;1;500,00;0;0;0;2,95",
            "08/03/2026;09/03/2026;VENDITA;ETF TEST;IE00TEST0001;V;2;EUR;100,00;1;200,00;0;0;0;2,95",
        ]
    )

    response = service.parse_public_portfolio_file(
        file_content=content,
        filename="fineco.csv",
        broker="fineco",
    )

    assert response.total_rows == 3
    assert response.valid_rows == 3
    assert response.error_rows == 0
    assert len(response.positions) == 1
    assert response.positions[0].identifier == "IE00TEST0001"
    assert response.positions[0].value == 1300.0
    assert "IE00TEST0001 1300.00" in response.raw_text


class _CommitRepo:
    def __init__(self, preview_data):
        self.preview_data = preview_data
        self.transactions = []

    def get_csv_import_batch(self, batch_id, user_id):
        return {"status": "pending", "portfolio_id": 1, "preview_data": self.preview_data}

    def create_transaction(self, payload, user_id):
        self.transactions.append(payload)

    def commit_csv_import_batch(self, batch_id, user_id):
        pass


def test_commit_batch_stores_timezone_aware_trade_at():
    repo = _CommitRepo(
        [
            {
                "row_number": 1,
                "valid": True,
                "trade_at": "2026-10-13T00:00:00",
                "isin": "IE00BK5BQT80",
                "side": "sell",
                "quantity": 1.0,
                "price": 153.89,
                "trade_currency": "EUR",
                "asset_id": 7,
            }
        ]
    )
    service = CsvImportService(repo)

    result = service.commit_batch(1, "user")

    assert result.errors == []
    assert result.committed_transactions == 1
    assert repo.transactions[0].trade_at.tzinfo is not None


# --- AI-assisted column mapping for formats without a dedicated parser ---


class _ImportRepo:
    """Minimal repository for parse_and_validate: keeps the saved batch and the commit calls."""

    def __init__(self, ai_key_enc: str = "") -> None:
        self.ai_key_enc = ai_key_enc
        self.batches: list[dict] = []
        self.committed_batches: list[int] = []

    def get_user_settings(self, user_id):
        return SimpleNamespace(copilot_provider="openai", copilot_model="")

    def get_user_copilot_api_key_enc(self, user_id):
        return self.ai_key_enc

    def search_assets(self, query):
        return []

    def create_csv_import_batch(self, **kwargs):
        self.batches.append(kwargs)
        return len(self.batches)

    def commit_csv_import_batch(self, batch_id, user_id):
        self.committed_batches.append(batch_id)


class _FakeProvider:
    """Stands in for call_provider_for_text: records the prompts, answers with a canned text."""

    def __init__(self, answer=None, error: Exception | None = None) -> None:
        self.answer = answer
        self.error = error
        self.calls: list[dict] = []

    def __call__(self, config, system_prompt, user_prompt, **kwargs):
        self.calls.append({"config": config, "system_prompt": system_prompt, "user_prompt": user_prompt, **kwargs})
        if self.error is not None:
            raise self.error
        return self.answer if isinstance(self.answer, str) else json.dumps(self.answer)


@pytest.fixture
def ai_service(monkeypatch):
    """A service whose user has an AI provider configured; the provider call is replaced per test."""
    repo = _ImportRepo()
    monkeypatch.setattr(
        "app.services.csv_service.resolve_copilot_config",
        lambda settings, **kwargs: CopilotConfig(provider="openai", model="test-model", api_key="k"),
    )

    def install(provider: _FakeProvider) -> CsvImportService:
        monkeypatch.setattr("app.services.csv_ai_mapping.call_provider_for_text", provider)
        return CsvImportService(repo)

    install.repo = repo
    return install


def _xlsx_bytes(rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


# Broker A: Italian export, semicolon separated, with a title block above the header.
BROKER_A_CSV = "\n".join(
    [
        "Conto: 000123;;;;;;;",
        "Movimenti titoli dal 01/03/2026 al 31/03/2026;;;;;;;",
        "",
        "Data operazione;Tipo;ISIN;Descrizione;Quantita;Importo unitario;Valuta;Commissioni",
        "02/03/2026;Acquisto;IE00B4L5Y983;ISHARES CORE MSCI WORLD;10;85,50;EUR;1,50",
        "09/03/2026;Acquisto;IE00BK5BQT80;VANGUARD FTSE ALL-WORLD;3;102,10;EUR;1,50",
        "16/03/2026;Vendita;IE00B4L5Y983;ISHARES CORE MSCI WORLD;4;88,00;EUR;1,50",
    ]
)
BROKER_A_MAPPING = {
    "trade_at": "Data operazione",
    "isin": "ISIN",
    "side": "Tipo",
    "quantity": "Quantita",
    "price": "Importo unitario",
    "trade_currency": "Valuta",
    "fees": "Commissioni",
    "name": "Descrizione",
    "date_format": "DD/MM/YYYY",
}

# Broker B: English XLSX export with native Excel dates and numbers, plus non-trade rows.
BROKER_B_ROWS = [
    ["Action", "Time", "ISIN", "Ticker", "Name", "No. of shares", "Price / share", "Currency (Price / share)", "Total"],
    ["Deposit", datetime(2026, 3, 1, 9, 0, 0), None, None, None, None, None, None, 1000.0],
    ["Market buy", datetime(2026, 3, 2, 10, 15, 30), "IE00B4L5Y983", "IWDA", "iShares Core MSCI World", 3.5, 98.12, "EUR", 343.42],
    ["Limit sell", datetime(2026, 3, 9, 15, 45, 0), "IE00B4L5Y983", "IWDA", "iShares Core MSCI World", 1.0, 99.5, "EUR", 99.5],
    ["Dividend (Ordinary)", datetime(2026, 3, 20, 8, 0, 0), "IE00B4L5Y983", "IWDA", "iShares Core MSCI World", 2.5, 0.31, "EUR", 0.78],
]
BROKER_B_MAPPING = {
    "trade_at": "Time",
    "isin": "ISIN",
    "side": "Action",
    "quantity": "No. of shares",
    "price": "Price / share",
    "trade_currency": "Currency (Price / share)",
    "fees": None,
    "name": "Name",
    "date_format": "YYYY-MM-DD",
}

BROKER_A_MISSING_COLUMNS_ERROR = "Colonne obbligatorie mancanti: operazione, segno, quantita, prezzo"


def test_unrecognised_csv_broker_gets_ai_mapping_in_the_preview(ai_service):
    provider = _FakeProvider(BROKER_A_MAPPING)
    service = ai_service(provider)

    # "fineco" is the UI default: its skip_rows profile must not eat the unknown file.
    response = service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv", broker="fineco")

    assert len(provider.calls) == 1
    assert response.column_mapping is not None
    assert response.column_mapping.source == "ai"
    assert response.column_mapping.mapping.model_dump() == BROKER_A_MAPPING
    assert response.column_mapping.columns == [
        "Data operazione", "Tipo", "ISIN", "Descrizione", "Quantita", "Importo unitario", "Valuta", "Commissioni",
    ]
    assert (response.total_rows, response.valid_rows, response.error_rows) == (3, 3, 0)

    first, second, third = response.rows
    assert first.trade_at == "2026-03-02T00:00:00"
    assert first.isin == "IE00B4L5Y983"
    assert first.side == "buy"
    assert first.quantity == 10
    assert first.price == 85.5
    assert first.fees == 1.5
    assert first.trade_currency == "EUR"
    assert first.titolo == "ISHARES CORE MSCI WORLD"
    assert second.side == "buy"
    assert third.side == "sell"

    # The preview is stored as a pending batch; nothing is committed without the user's confirmation.
    assert len(service.repo.batches) == 1
    assert service.repo.batches[0]["valid_rows"] == 3
    assert service.repo.committed_batches == []


@pytest.mark.parametrize("broker", ["generic", "fineco"])
def test_unrecognised_xlsx_broker_gets_ai_mapping_in_the_preview(ai_service, broker):
    provider = _FakeProvider(BROKER_B_MAPPING)
    service = ai_service(provider)

    response = service.parse_and_validate(
        1, "user", file_bytes=_xlsx_bytes(BROKER_B_ROWS), filename="trading.xlsx", broker=broker
    )

    assert len(provider.calls) == 1
    assert response.column_mapping is not None
    assert response.column_mapping.source == "ai"
    assert response.column_mapping.mapping.fees is None
    assert response.total_rows == 4

    deposit, buy, sell, dividend = response.rows
    assert buy.valid and sell.valid
    assert buy.trade_at == "2026-03-02T10:15:30"
    assert buy.side == "buy"
    assert buy.quantity == 3.5
    assert buy.price == 98.12
    assert buy.trade_currency == "EUR"
    assert sell.side == "sell"
    # Rows that are not buys or sells are flagged in the preview, never imported silently.
    assert not deposit.valid
    assert not dividend.valid
    assert any("segno non valido" in error for error in dividend.errors)
    assert response.valid_rows == 2
    assert service.repo.committed_batches == []


def test_ai_receives_only_headers_and_at_most_five_sample_rows(ai_service):
    header = "Data;Tipo;ISIN;Quantita;Prezzo;Note"
    body = [f"0{day}/03/2026;Acquisto;IE00B4L5Y983;{day};10,00;riga-segreta-{day}" for day in range(1, 10)]
    body[0] = body[0].replace("riga-segreta-1", "x" * 200)
    provider = _FakeProvider(
        {"trade_at": "Data", "isin": "ISIN", "side": "Tipo", "quantity": "Quantita", "price": "Prezzo"}
    )
    service = ai_service(provider)

    response = service.parse_and_validate(1, "user", file_content="\n".join([header, *body]), filename="x.csv")

    assert response.total_rows == 9
    (call,) = provider.calls
    sent = json.loads(call["user_prompt"])
    assert sent["columns"] == header.split(";")
    assert len(sent["sample_rows"]) == 5
    assert set(sent) == {"columns", "sample_rows"}
    # Rows beyond the fifth never leave the server; long cells are cut.
    assert "riga-segreta-6" not in call["user_prompt"]
    assert "riga-segreta-9" not in call["user_prompt"]
    assert max(len(cell) for row in sent["sample_rows"] for cell in row.values()) <= 60
    # The user's key (or the server key) is the one resolved by the copilot configuration.
    assert call["config"].model == "test-model"
    # A stuck provider must not hang the import: the call is bounded in time and in size.
    assert call["timeout"] > 0
    assert call["max_tokens"] >= 200


def test_ai_answer_wrapped_in_text_and_code_fence_is_accepted(ai_service):
    answer = "Ecco la mappatura:\n```json\n" + json.dumps(BROKER_A_MAPPING) + "\n```\n"
    service = ai_service(_FakeProvider(answer))

    response = service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv")

    assert response.column_mapping is not None
    assert response.valid_rows == 3


GENERIC_CSV = "\n".join(
    [
        "operazione;isin;segno;quantita;prezzo",
        "02/03/2026;IE00B4L5Y983;A;10;85,50",
    ]
)
FINECO_CSV = "\n".join(
    [
        "Dossier n.: 00000000",
        "RISULTATO RICERCA MOVIMENTI TITOLI",
        ";".join(REDUCED_FINECO_HEADER),
        "13/10/2026;15/10/2026;Compravendita titoli;ETF;IE00B4L60045;A;375;EUR;105,23;1;39461,5;19",
    ]
)


@pytest.mark.parametrize(
    "kind,payload",
    [
        ("csv", GENERIC_CSV),
        ("csv", FINECO_CSV),
        ("xlsx", _xlsx_bytes([["Operazione", "Isin", "Segno", "Quantita", "Prezzo"], ["02/03/2026", "IE00B4L5Y983", "A", 10, 85.5]])),
    ],
    ids=["generic-csv", "fineco-csv", "generic-xlsx"],
)
def test_supported_formats_never_go_through_the_ai(ai_service, kind, payload):
    provider = _FakeProvider(error=AssertionError("the AI must not be called for supported formats"))
    service = ai_service(provider)
    kwargs = {"file_content": payload, "filename": "f.csv"} if kind == "csv" else {"file_bytes": payload, "filename": "f.xlsx"}

    response = service.parse_and_validate(1, "user", broker="fineco" if payload is FINECO_CSV else "generic", **kwargs)

    assert provider.calls == []
    assert response.column_mapping is None
    assert response.valid_rows == 1


def test_without_ai_the_original_error_is_unchanged(monkeypatch):
    monkeypatch.setattr("app.services.csv_service.resolve_copilot_config", lambda settings, **kwargs: None)
    provider = _FakeProvider(error=AssertionError("no provider is configured"))
    monkeypatch.setattr("app.services.csv_ai_mapping.call_provider_for_text", provider)
    repo = _ImportRepo()
    service = CsvImportService(repo)

    with pytest.raises(ValueError) as csv_error:
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv", broker="generic")
    with pytest.raises(ValueError) as xlsx_error:
        service.parse_and_validate(
            1, "user", file_bytes=_xlsx_bytes(BROKER_B_ROWS), filename="trading.xlsx", broker="generic"
        )

    # The header block above the real header makes the CSV fail before the column check.
    assert str(csv_error.value) == "Colonne obbligatorie mancanti: operazione, isin, segno, quantita, prezzo"
    assert str(xlsx_error.value) == "Colonne obbligatorie mancanti: operazione, segno, quantita, prezzo"
    assert provider.calls == []
    assert repo.batches == []


def test_unsupported_provider_counts_as_ai_not_available(monkeypatch):
    monkeypatch.setattr(
        "app.services.csv_service.resolve_copilot_config",
        lambda settings, **kwargs: CopilotConfig(provider="sconosciuto", model="m", api_key="k"),
    )
    repo = _ImportRepo()
    service = CsvImportService(repo)

    with pytest.raises(ValueError, match="Colonne obbligatorie mancanti"):
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv")

    assert repo.batches == []


def test_ai_configuration_lookup_failure_keeps_the_original_error(monkeypatch):
    class _BrokenRepo(_ImportRepo):
        def get_user_settings(self, user_id):
            raise RuntimeError("db down")

    service = CsvImportService(_BrokenRepo())

    with pytest.raises(ValueError, match="Colonne obbligatorie mancanti"):
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv")


@pytest.mark.parametrize(
    "answer,error",
    [
        (None, RuntimeError("provider timeout")),
        ("Non riesco a mappare queste colonne.", None),
        ("{not json}", None),
        ({**BROKER_A_MAPPING, "controvalore": "Commissioni"}, None),  # unknown target field
        ({**BROKER_A_MAPPING, "isin": "Codice titolo"}, None),  # column not in the file
        ({**BROKER_A_MAPPING, "isin": None}, None),  # required field missing
        ({**BROKER_A_MAPPING, "date_format": "dd/mm/aa"}, None),  # format outside the whitelist
        ({**BROKER_A_MAPPING, "price": "Quantita"}, None),  # same column for two fields
        ({**BROKER_A_MAPPING, "trade_at": "Descrizione"}, None),  # dates not parseable
        ({**BROKER_A_MAPPING, "trade_at": "Data operazione", "date_format": "YYYY-MM-DD"}, None),  # wrong date format
        ({**BROKER_A_MAPPING, "quantity": "Descrizione"}, None),  # amounts not parseable
        ({**BROKER_A_MAPPING, "isin": "Tipo"}, None),  # not ISIN codes
        ({**BROKER_A_MAPPING, "side": "Descrizione"}, None),  # no buy/sell values
        ({**BROKER_A_MAPPING, "trade_currency": "Descrizione"}, None),  # not currency codes
        ({**BROKER_A_MAPPING, "fees": "Descrizione"}, None),  # fees not parseable
    ],
)
def test_failed_or_invalid_ai_proposal_keeps_the_original_error(ai_service, answer, error):
    service = ai_service(_FakeProvider(answer, error))

    with pytest.raises(ValueError) as exc_info:
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="directa.csv")

    assert not isinstance(exc_info.value, ColumnMappingError)
    assert str(exc_info.value) == "Colonne obbligatorie mancanti: operazione, isin, segno, quantita, prezzo"
    assert service.repo.batches == []


def test_user_corrected_mapping_is_applied_without_calling_the_ai(ai_service):
    provider = _FakeProvider(error=AssertionError("the AI must not be called when the user sends a mapping"))
    service = ai_service(provider)
    corrected = CsvColumnMapping(**{**BROKER_A_MAPPING, "fees": None, "trade_currency": None})

    response = service.parse_and_validate(
        1, "user", file_content=BROKER_A_CSV, filename="directa.csv", column_mapping=corrected
    )

    assert provider.calls == []
    assert response.column_mapping is not None
    assert response.column_mapping.source == "user"
    assert response.column_mapping.mapping == corrected
    assert response.valid_rows == 3
    assert response.rows[0].fees is None
    assert response.rows[0].trade_currency is None
    assert service.repo.committed_batches == []


def test_user_mapping_that_does_not_fit_the_file_is_rejected(ai_service):
    service = ai_service(_FakeProvider(error=AssertionError("unexpected AI call")))
    wrong = CsvColumnMapping(**{**BROKER_A_MAPPING, "quantity": "Descrizione", "name": None})
    reused = CsvColumnMapping(**{**BROKER_A_MAPPING, "price": "Quantita"})
    missing_column = CsvColumnMapping(**{**BROKER_A_MAPPING, "isin": "Codice"})

    with pytest.raises(ColumnMappingError, match="'Descrizione'.*quantita"):
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="d.csv", column_mapping=wrong)
    with pytest.raises(ColumnMappingError, match="'Quantita' non puo' essere assegnata"):
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="d.csv", column_mapping=reused)
    with pytest.raises(ColumnMappingError, match="'Codice'.*non e' presente"):
        service.parse_and_validate(1, "user", file_content=BROKER_A_CSV, filename="d.csv", column_mapping=missing_column)

    assert service.repo.batches == []


def test_row_with_a_date_outside_the_declared_format_is_flagged_not_reinterpreted(ai_service):
    # Dates are checked on the sample (first five rows); later rows are checked one by one in the preview.
    content = "\n".join(
        [
            "Data;Tipo;ISIN;Quantita;Prezzo",
            *["03/04/2026;Buy;IE00B4L5Y983;1;10,00"] * 5,
            "13/01/2026;Buy;IE00B4L5Y983;1;10,00",
        ]
    )
    mapping = CsvColumnMapping(trade_at="Data", isin="ISIN", side="Tipo", quantity="Quantita", price="Prezzo", date_format="MM/DD/YYYY")
    service = ai_service(_FakeProvider(error=AssertionError("unexpected AI call")))

    response = service.parse_and_validate(1, "user", file_content=content, filename="us.csv", column_mapping=mapping)

    assert response.rows[0].trade_at == "2026-03-04T00:00:00"
    assert response.valid_rows == 5
    assert not response.rows[5].valid
    assert "Formato data non riconosciuto" in response.rows[5].errors[0]


def test_sample_validation_tolerates_non_trade_rows_with_blank_fields(ai_service):
    # Broker B starts with a deposit row without ISIN, quantity or price: blank cells are not errors.
    service = ai_service(_FakeProvider(BROKER_B_MAPPING))

    response = service.parse_and_validate(1, "user", file_bytes=_xlsx_bytes(BROKER_B_ROWS), filename="t.xlsx")

    assert response.column_mapping is not None


def test_read_mapping_table_finds_header_below_a_preamble_and_the_right_delimiter():
    columns, rows = _read_mapping_table(BROKER_A_CSV, "a.csv", None)

    assert columns[:3] == ["Data operazione", "Tipo", "ISIN"]
    assert len(rows) == 3

    comma_csv = 'Titolo report\nData,Tipo,Importo\n01/03/2026,Buy,"1,234.50"\n'
    columns, rows = _read_mapping_table(comma_csv, "b.csv", None)
    assert columns == ["Data", "Tipo", "Importo"]
    assert rows == [{"Data": "01/03/2026", "Tipo": "Buy", "Importo": "1,234.50"}]

    tab_csv = "Data\tTipo\tQta\n01/03/2026\tBuy\t3\n"
    assert _read_mapping_table(tab_csv, "c.csv", None)[0] == ["Data", "Tipo", "Qta"]


def test_read_mapping_table_rejects_files_without_data_rows():
    with pytest.raises(ValueError):
        _read_mapping_table("Data;Tipo;ISIN\n", "a.csv", None)
    with pytest.raises(ValueError):
        _read_mapping_table("", "a.csv", None)


@pytest.mark.parametrize(
    "value,date_format,expected",
    [
        ("02/03/2026", "DD/MM/YYYY", datetime(2026, 3, 2)),
        ("02/03/2026", "MM/DD/YYYY", datetime(2026, 2, 3)),
        ("02.03.2026 14:30", "DD.MM.YYYY", datetime(2026, 3, 2, 14, 30)),
        ("2026-03-02T10:15:30Z", "YYYY-MM-DD", datetime(2026, 3, 2, 10, 15, 30)),
        ("2026-03-02 10:15:30.123456", "DD/MM/YYYY", datetime(2026, 3, 2, 10, 15, 30)),
        ("2026-03-02 10:15:30+02:00", "YYYY-MM-DD", datetime(2026, 3, 2, 10, 15, 30)),
        ("02-03-2026", "DD-MM-YYYY", datetime(2026, 3, 2)),
        ("2026/03/02", "YYYY/MM/DD", datetime(2026, 3, 2)),
        ("13/01/2026", "MM/DD/YYYY", None),
        ("domani", "DD/MM/YYYY", None),
    ],
)
def test_parse_mapped_datetime(value, date_format, expected):
    assert _parse_mapped_datetime(value, date_format) == expected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("A", "A"), ("Acquisto", "A"), ("Market buy", "A"), ("BUY", "A"), ("Compra", "A"),
        ("V", "V"), ("Vendita", "V"), ("Limit sell", "V"), ("sold", "V"),
        ("Dividendo", None), ("Dividend (Ordinary)", None), ("Deposit", None), ("Compravendita", None), ("", None),
    ],
)
def test_side_from_value(value, expected):
    assert _side_from_value(value) == expected


def test_call_provider_for_text_forwards_max_tokens_and_timeout_and_keeps_its_defaults(monkeypatch):
    import anthropic
    import openai
    from google import genai

    from app.services.instant_portfolio_analyzer import call_provider_for_text

    seen: dict[str, dict] = {}

    class FakeOpenAI:
        def __init__(self, **kwargs):
            seen["openai_init"] = kwargs
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

        def _create(self, **kwargs):
            seen["openai_call"] = kwargs
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="ok"))])

    class FakeAnthropic:
        def __init__(self, **kwargs):
            seen["anthropic_init"] = kwargs
            self.messages = SimpleNamespace(create=self._create)

        def _create(self, **kwargs):
            seen["anthropic_call"] = kwargs
            return SimpleNamespace(content=[SimpleNamespace(type="text", text="ok")])

    class FakeGenai:
        def __init__(self, **kwargs):
            seen["gemini_init"] = kwargs
            self.models = SimpleNamespace(generate_content=self._generate)

        def _generate(self, **kwargs):
            seen["gemini_call"] = kwargs
            return SimpleNamespace(text="ok")

    monkeypatch.setattr(openai, "OpenAI", FakeOpenAI)
    monkeypatch.setattr(anthropic, "Anthropic", FakeAnthropic)
    monkeypatch.setattr(genai, "Client", FakeGenai)

    for provider in ("openai", "anthropic", "gemini"):
        config = CopilotConfig(provider=provider, model="m", api_key="k")
        assert call_provider_for_text(config, "sys", "user") == "ok"
        assert call_provider_for_text(config, "sys", "user", max_tokens=400, timeout=30.0) == "ok"

    # Without the new arguments the calls are the same as before: 220 tokens, SDK default timeout.
    assert seen["openai_call"]["max_tokens"] == 400
    assert seen["anthropic_call"]["max_tokens"] == 400
    assert seen["gemini_call"]["config"] == {"max_output_tokens": 400}
    assert seen["openai_init"]["timeout"] == 30.0
    assert seen["anthropic_init"]["timeout"] == 30.0
    assert seen["gemini_init"]["http_options"] == {"timeout": 30000}

    config = CopilotConfig(provider="openai", model="m", api_key="k")
    call_provider_for_text(config, "sys", "user")
    assert seen["openai_call"]["max_tokens"] == 220
    assert "timeout" not in seen["openai_init"]
    config = CopilotConfig(provider="anthropic", model="m", api_key="k")
    call_provider_for_text(config, "sys", "user")
    assert seen["anthropic_call"]["max_tokens"] == 220
    assert "timeout" not in seen["anthropic_init"]
    config = CopilotConfig(provider="gemini", model="m", api_key="k")
    call_provider_for_text(config, "sys", "user")
    assert seen["gemini_call"]["config"] == {"max_output_tokens": 220}
    assert seen["gemini_init"]["http_options"] is None
